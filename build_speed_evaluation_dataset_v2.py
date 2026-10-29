import collections
import random

import datasets
import numpy as np
import terminaltables
import tqdm
import json
import traceback
import re
import os
import openpyxl
from openpyxl.utils import range_boundaries
import torch
import transformers
import unicodedata
import matplotlib.pyplot as plt
import gzip
import json
import requests
import langid  # pip install langid


def is_russian_text(text: str) -> bool:
    """Perform fast check if provided text looks Russian (cyrillic, strictly speaking)."""

    # Set of Russian Cyrillic letters (including Ё/ё, which are outside the А-Я range)
    russian_letters = set(
        "абвгдеёжзийклмнопрстуфхцчшщъыьэюя"
        "АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ"
    )

    # Keep only characters that are neither whitespace nor punctuation
    relevant = [
        ch for ch in text
        if not ch.isspace() and not unicodedata.category(ch).startswith("P")
    ]

    # Nothing meaningful to analyze
    if not relevant:
        return False

    russian_count = sum(1 for ch in relevant if ch in russian_letters)
    return russian_count / len(relevant) > 0.5


def is_han(ch: str) -> bool:
    cp = ord(ch)
    return (
        0x3400 <= cp <= 0x4DBF      # CJK Extension A
        or 0x4E00 <= cp <= 0x9FFF   # CJK Unified Ideographs
        or 0xF900 <= cp <= 0xFAFF   # CJK Compatibility Ideographs
        or 0x20000 <= cp <= 0x2FA1F # CJK Extensions B-F and related ranges
    )


def is_english_text(text: str) -> bool:
    """Return True when a lightweight language detector identifies text as English.

    `langid` is a small, pure-Python dependency and does not require downloading
    a separate language-identification model. A 4,000-character prefix is enough
    for this dataset's prompt-length limit and avoids needless work on long inputs.
    """
    if not text or not text.strip():
        return False

    if re.search("[абвгдеёжзийклмнопрстуфхцчшщъыьэюя]", text, flags=re.I):
        return False

    if any(is_han(c) for c in text):
        return False

    if re.search("traduire", text, flags=re.I):
        return False

    # Language identification is noisy for very short strings and code-only text;
    # the dataset's prompt-length filter is applied before this check.
    detected_language, _score = langid.classify(text[:4000])
    return detected_language == "en"



def is_chinese_text(text: str) -> bool:
    """Return True when Han characters make up most non-space, non-punctuation characters."""

    relevant = [
        ch for ch in text
        if not ch.isspace() and not unicodedata.category(ch).startswith("P")
    ]
    if not relevant:
        return False

    han_count = sum(1 for ch in relevant if is_han(ch))
    return han_count / len(relevant) > 0.5


def is_general_qa_sample(prompt: str, response: str) -> bool:
    if "```" in response:
        return False
    return True


def check_prompt_length(prompt: str) -> bool:
    if prompt is None or len(prompt) < 50 or len(prompt) > 4000:
        return False

    prompt_token_ids = tokenizer.apply_chat_template([{"role": "user", "content": prompt}],
                                                    add_generation_prompt=True,
                                                    enable_thinking=False,
                                                    return_dict=True)["input_ids"]
    if len(prompt_token_ids) > MAX_PROMPT_TOKENS:
        return False

    return True


def check_chinese_prompt_length(prompt: str) -> bool:
    if prompt is None or len(prompt) < 10 or len(prompt) > 4000:
        return False

    prompt_token_ids = tokenizer.apply_chat_template([{"role": "user", "content": prompt}],
                                                    add_generation_prompt=True,
                                                    enable_thinking=False,
                                                    return_dict=True)["input_ids"]
    if len(prompt_token_ids) > MAX_PROMPT_TOKENS:
        return False

    return True


def get_cell_range(file_path, cell_range, sheet_name=None):
    """Get values from a range of cells (e.g., 'A1:C5')"""
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)

        if sheet_name:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active

        # Parse the range
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)

        values = []
        for row in range(min_row, max_row + 1):
            row_values = []
            for col in range(min_col, max_col + 1):
                row_values.append(sheet.cell(row=row, column=col).value)
            values.append(row_values)

        return values

    except Exception as e:
        print(f"Error: {e}")
        return None


def build_lorugec():
    # LORuGEC
    samples = []
    fp = os.path.join("/home/inkoziev/polygon/text_generator/data/speller/LORuGEC/LORuGEC.xlsx")
    pairs = get_cell_range(fp, 'D2:G961')
    for i, (rule, _, corrupted, fixed) in enumerate(pairs):
        if rule in ("Punctuation",):
            continue

        corrupted = re.sub(r"[ ]{1,}([.,?!])", r"\1", corrupted)
        fixed = re.sub(r"[ ]{1,}([.,?!])", r"\1", fixed)
        if corrupted != fixed:
            # есть дефект
            samples.append({'text': corrupted, 'label': 1, 'domain': "LORuGEC"})

    return samples


def build_jfleg():
    raw = list(datasets.load_dataset('jhu-clsp/jfleg', split='test'))

    def convert(raw_samples):
        grammatical = []
        ungrammatical = []

        for sample in raw_samples:
            corrupted = sample['sentence']

            for corrected in sample['corrections']:
                if corrupted.lower() != corrected.lower():
                    grammatical.append({
                        'text': corrected,
                        'label': 0,
                        'domain': 'jfleg',
                    })

                    ungrammatical.append({
                        'text': corrupted,
                        'label': 1,
                        'domain': 'jfleg',
                    })

        return grammatical + ungrammatical

    return convert(raw)


def stop_token_ids(model, tokenizer) -> set:
    """Every id that ends a chat response: the generation config's EOS ids (what generate() stops on), the tokenizer's EOS
    and <|im_end|>. Counting up to the first of them also ignores the pad tokens that follow a finished row in a batch."""
    ids = set()
    for x in (getattr(model.generation_config, "eos_token_id", None), tokenizer.eos_token_id,
              tokenizer.convert_tokens_to_ids("<|im_end|>")):
        if x is None:
            continue
        ids.update(x if isinstance(x, (list, tuple)) else [x])
    return {int(i) for i in ids if i is not None and i >= 0}


def stream_jsonl_gz(url):
    """Yield parsed JSON objects from a remote .jsonl.gz file without
    downloading it completely."""
    with requests.get(url, stream=True, timeout=30) as resp:
        resp.raise_for_status()
        # GzipFile reads the raw response stream and decompresses on the fly
        with gzip.GzipFile(fileobj=resp.raw) as gz:
            for line in gz:
                line = line.strip()
                if not line:
                    continue
                try:
                    yield json.loads(line)
                except json.JSONDecodeError:
                    continue



@torch.inference_mode()
def measure_response_lengths(model, tokenizer, prompts, batch_size: int, max_new_tokens: int, desc: str = ""):
    """Measure greedy response lengths, splitting OOM batches without recursive retries.

    Prompts that still OOM when generated individually are skipped rather than
    aborting the entire dataset build. Their response length is unknown, so they
    are excluded from the returned list.
    """
    stop_ids = stop_token_ids(model, tokenizer)
    pad_id = tokenizer.pad_token_id
    if pad_id is None:
        pad_id = next(iter(stop_ids), tokenizer.eos_token_id)
    if pad_id is None:
        raise ValueError("Tokenizer must define a pad token or an EOS/stop token ID")

    encoded = [
        tokenizer.apply_chat_template(
            [{"role": "user", "content": p}],
            add_generation_prompt=True,
            enable_thinking=False,
            return_dict=True,
        )["input_ids"]
        for p in prompts
    ]
    order = sorted(range(len(prompts)), key=lambda i: -len(encoded[i]))
    response_len = [None] * len(prompts)

    def try_generate(idx):
        """Return generated token IDs on CPU, or None on CUDA OOM.

        Keep this attempt in a separate stack frame. In particular, do not
        recursively retry from inside an except block: the OOM traceback can
        retain references to tensors from the failed generation while retries
        are running, preventing those allocations from being released.
        """
        input_ids = attention_mask = out = None
        try:
            L = max(len(encoded[i]) for i in idx)
            input_ids = torch.full((len(idx), L), pad_id, dtype=torch.long)
            attention_mask = torch.zeros((len(idx), L), dtype=torch.long)
            for r, i in enumerate(idx):
                ids = encoded[i]
                input_ids[r, L - len(ids):] = torch.tensor(ids, dtype=torch.long)
                attention_mask[r, L - len(ids):] = 1

            out = model.generate(
                input_ids=input_ids.to(model.device),
                attention_mask=attention_mask.to(model.device),
                do_sample=False,
                temperature=None,
                top_p=None,
                top_k=None,
                max_new_tokens=max_new_tokens,
                pad_token_id=pad_id,
            )
            # Copy generated tokens to CPU before releasing the GPU output.
            return out[:, L:].cpu().tolist()
        except torch.cuda.OutOfMemoryError:
            return None
        finally:
            del out, input_ids, attention_mask

    # Use an explicit work queue rather than recursive calls. Failed batch
    # frames are gone before smaller batches are attempted.
    pending = list(reversed([
        order[k:k + batch_size] for k in range(0, len(order), batch_size)
    ]))
    pbar = tqdm.tqdm(total=len(prompts), desc=desc)
    try:
        while pending:
            idx = pending.pop()
            generated = try_generate(idx)
            if generated is None:
                # The failed attempt's frame/traceback has now unwound, so the
                # allocator can release any newly-unused cached CUDA blocks.
                if torch.cuda.is_available():
                    torch.cuda.empty_cache()
                if len(idx) > 1:
                    mid = len(idx) // 2
                    # Push second half first so the first half is processed next.
                    pending.append(idx[mid:])
                    pending.append(idx[:mid])
                else:
                    i = idx[0]
                    print(
                        f"WARNING: skipping prompt {i} in {desc!r}: CUDA OOM "
                        f"even with batch_size=1 (prompt_tokens={len(encoded[i])})."
                    )
                    pbar.update(1)
                continue

            for r, i in enumerate(idx):
                n = next(
                    (k + 1 for k, token_id in enumerate(generated[r]) if token_id in stop_ids),
                    len(generated[r]),
                )
                response_len[i] = min(n, max_new_tokens)
                pbar.update(1)
    finally:
        pbar.close()

    # OOM-skipped prompts have unknown lengths and must not be ranked as if
    # they had zero-length responses.
    return [
        (prompt, len(ids), n)
        for prompt, ids, n in zip(prompts, encoded, response_len)
        if n is not None
    ]


def plot_grouped_length_histograms(
    domain: str,
    data: dict,
    bins: int = 30,
    density: bool = False,
    gap: float = 0.2,        # fraction of bin width left empty between bins
    figsize=(10, 6),
    alpha: float = 0.75,
):
    """
    Grouped histograms of response lengths: shared bins, side-by-side bars,
    with a visible gap between adjacent bins.

    Parameters
    ----------
    domain : str
        Domain name
    data : dict[str, list[float]]
        Dictionary {language_code: [response_lengths]}.
    bins : int
        Number of bins (shared across all languages).
    density : bool
        Normalize each language's histogram (useful for unequal sample sizes).
    gap : float
        Fraction of the bin width left empty between bins (0..1).
        E.g. 0.2 means bars occupy 80% of the bin, 20% is a gap.
    figsize : tuple
        Figure size.
    alpha : float
        Bar opacity.
    """
    if not 0.0 <= gap < 1.0:
        raise ValueError("gap must be in [0, 1)")

    data = {lang: np.asarray(vals, dtype=float)
            for lang, vals in data.items() if len(vals) > 0}
    if not data:
        raise ValueError("No data to visualize")

    # Shared bin edges — identical for every language
    all_values = np.concatenate(list(data.values()))
    bin_edges = np.linspace(all_values.min(), all_values.max(), bins + 1)
    bin_width = bin_edges[1] - bin_edges[0]
    lefts = bin_edges[:-1]

    langs = list(data.keys())
    n_langs = len(langs)

    # Width actually occupied by bars inside one bin (rest = gap)
    group_width = bin_width * (1.0 - gap)
    bar_width = group_width / n_langs

    # Offset of the group's left edge from the bin's left edge
    group_offset = (bin_width - group_width) / 2.0

    fig, ax = plt.subplots(figsize=figsize)

    for i, lang in enumerate(langs):
        counts, _ = np.histogram(data[lang], bins=bin_edges, density=density)

        # Bars sit side-by-side, starting at group_offset inside the bin
        positions = lefts + group_offset + (i + 0.5) * bar_width

        ax.bar(
            positions,
            counts,
            width=bar_width,
            alpha=alpha,
            label=lang,
            edgecolor="black",
            linewidth=0.4,
        )

    ax.set_xlabel("Response length, tokens")
    ax.set_ylabel("Density" if density else "Count")
    ax.set_title(f"Distribution of LLM response lengths in domain {domain}")
    ax.legend(title="Language")
    ax.grid(True, alpha=0.3, axis="y")
    plt.tight_layout()

    return fig, ax


if __name__ == "__main__":
    N_REPEATS = 100             # Number of prompts to keep per domain
    N_PREFETCH = 10 * N_REPEATS  # Number of prompts to sample from domain pool before ranking and selecting.
    GEN_BATCH_SIZE = 32         # prompts per generate() call; an OOM batch is split automatically. This setting is tuned for my RTX 3090 with 24 Gb of VRAM.
    MAX_NEW_TOKENS = 2048       # response-length probe; responses reaching it count as MAX_NEW_TOKENS
    MAX_PROMPT_TOKENS = MAX_NEW_TOKENS//2

    prompts = collections.defaultdict(list)

    shuffle_seed = 31415926

    rng = random.Random(shuffle_seed)

    model_id = "Qwen/Qwen3-0.6B"

    # ---------------------------------------------------------------------------

    tokenizer = transformers.AutoTokenizer.from_pretrained(model_id, padding_side="left")

    # Translation prompt templates.
    translation_templates = [
        # Neutral and direct
        "Translate this text into English: {}",
        "Translate the following text into English: {}",
        "Translate the text below from Russian to English: {}",
        "Translate a text into English: {}",
        "English translation: {}",

        # Alternative formulations
        'Given the text "{}", translate it into English.',
        'Translate the following Russian text into English: "{}"',
        'What is the English translation of the following text? {}',
        'Please provide an English translation of this text: {}',
        'Convert the following text from Russian into English: {}',

        # Polite and conversational
        "Could you translate the following text into English, please? {}",
        "Would you mind translating this text into English? {}",
        "Could you please provide an English translation of the following? {}",
        "I'd appreciate it if you could translate this into English: {}",
        "Please help me translate the following Russian text into English: {}",

        # Informal
        "Can you translate this into English? {}",
        "How would you say this in English? {}",
        "What does this sound like in English? {}",
        "Put this Russian text into English, please: {}",
        "Let's translate this into English: {}",

        # Literary and bookish
        "Render the following Russian passage in English: {}",
        "Provide an English rendition of the following text: {}",
        "Translate the passage below into clear, natural English: {}",
        "Please render the meaning of this Russian text in English: {}",
        "How might this passage be expressed in English? {}",
    ]

    # Russian grammar-error-correction prompt templates.
    gec_ru_templates = [
        "Исправь грамматические, пунктуационные и орфографические ошибки в тексте: {}",
        "Исправь ошибки в грамматике, пунктуации и орфографии следующего текста: {}",
        "Пожалуйста, исправь грамматические, пунктуационные и орфографические ошибки: {}",
        "Отредактируй текст, исправив грамматические, пунктуационные и орфографические ошибки: {}",
        "Проверь этот текст и исправь ошибки в грамматике, пунктуации и написании слов: {}",
        "Исправь все ошибки в следующем тексте, включая грамматику, пунктуацию и орфографию: {}",
        "Приведи текст в грамматически, пунктуационно и орфографически правильный вид: {}",
        "Найди и исправь грамматические, пунктуационные и орфографические ошибки в тексте: {}",
        "Устрани ошибки грамматики, пунктуации и орфографии в этом тексте: {}",
        "Отредактируй следующий текст, исправив ошибки, но сохранив его смысл: {}",
        "Исправь грамматику, расставь знаки препинания и исправь орфографию в тексте: {}",
        "Пожалуйста, вычитай текст и исправь ошибки в грамматике, пунктуации и орфографии: {}",
        "Сделай текст грамотным: исправь ошибки в грамматике, пунктуации и орфографии. Текст: {}",
        "Перепроверь этот текст и внеси необходимые исправления в грамматику, пунктуацию и орфографию: {}",
        "Отредактируй приведённый ниже текст, устранив грамматические и орфографические ошибки и ошибки пунктуации: {}",
        "Исправь языковые ошибки в тексте — грамматические, пунктуационные и орфографические: {}",
        "Помоги исправить ошибки в этом тексте, не меняя его исходного смысла: {}",
        "Пожалуйста, приведи следующий текст в соответствие с нормами русской грамматики, пунктуации и орфографии: {}",
        "Выполни корректуру текста и исправь ошибки в грамматике, пунктуации и написании слов: {}",
        "Улучши корректность текста, исправив грамматические, пунктуационные и орфографические ошибки: {}",
    ]

    # English grammar-error-correction prompt templates.
    gec_en_templates = [
        "Correct grammar, punctuation, and spelling errors in the text: {}",
        "Fix the grammatical, punctuation, and spelling mistakes in this text: {}",
        "Please correct any grammar, punctuation, or spelling errors in the following text: {}",
        "Proofread the text and correct its grammar, punctuation, and spelling: {}",
        "Edit the following text to fix grammatical, punctuation, and spelling errors: {}",
        "Find and correct all grammar, punctuation, and spelling mistakes in this text: {}",
        "Please proofread and correct the following: {}",
        "Correct the errors in grammar, spelling, and punctuation in the passage below: {}",
        "Review this text and fix any issues with grammar, punctuation, or spelling: {}",
        "Make the following text grammatically correct, with proper spelling and punctuation: {}",
        "Revise the text to correct its grammar, spelling, and punctuation while preserving its meaning: {}",
        "Clean up the grammar, punctuation, and spelling in this passage: {}",
        "Please edit this text for grammatical correctness, spelling, and punctuation: {}",
        "Check the following text for errors and correct its grammar, punctuation, and spelling: {}",
        "Polish this text by fixing any grammatical, spelling, or punctuation mistakes: {}",
        "Correct the language errors in the following text, including grammar, spelling, and punctuation: {}",
        "Could you proofread this text and fix its grammar, punctuation, and spelling? {}",
        "Please make any necessary corrections to the grammar, punctuation, and spelling of this text: {}",
        "Rewrite the text with its grammar, punctuation, and spelling errors corrected: {}",
        "Ensure that the following text uses correct grammar, spelling, and punctuation: {}",
    ]

    # Creative English story prompts. Keep the "about" and "where" pools
    # separate because their placeholders have different grammatical roles.
    creative_en_story_about_templates = [
        # Neutral and direct
        "Write a story inspired by a post on r/AskReddit about {}.",
        "Create a story based on an r/AskReddit discussion about {}.",
        "Write a fictional story about {}, as if inspired by r/AskReddit.",
        "Turn this r/AskReddit topic into a story: {}.",
        "Write a short story about a topic discussed on r/AskReddit: {}.",

        # Conversational
        "Imagine someone on r/AskReddit asks about {}. Write a story inspired by it.",
        "Someone on r/AskReddit brings up {}. Turn the topic into a story.",
        "Take the idea of {} from r/AskReddit and turn it into a compelling story.",
        "Write a story based on the kinds of personal experiences people share on r/AskReddit about {}.",
        "Imagine a discussion on r/AskReddit about {} and write a story inspired by it.",

        # Creative and literary
        "Craft a fictional narrative inspired by an r/AskReddit post about {}.",
        "Use the r/AskReddit topic of {} as the starting point for a story.",
        "Develop a narrative around the subject of {}, inspired by r/AskReddit.",
        "Let the r/AskReddit theme of {} inspire an engaging short story.",
        "Spin a compelling tale inspired by an r/AskReddit discussion about {}.",

        # Informal and playful
        "Here's an r/AskReddit-style topic: {}. Turn it into a story.",
        "Take {}—the sort of thing you'd see on r/AskReddit—and write a story about it.",
        "Write a fun or intriguing story inspired by an r/AskReddit thread about {}.",
        "Picture a post on r/AskReddit about {}. What story could come out of it?",
        "Turn an r/AskReddit discussion about {} into a fictional tale.",
    ]

    creative_en_story_where_templates = [
        # Neutral and direct
        "Write a story inspired by an r/AskReddit post where {}.",
        "Create a story based on an r/AskReddit scenario where {}.",
        "Write a fictional story set in a situation where {}.",
        "Tell a story about a scenario in which {}.",
        "Write a short story inspired by the idea of {} happening.",

        # Conversational
        "Imagine an r/AskReddit post describing a situation where {}. Write a story about it.",
        "Suppose someone on r/AskReddit describes a situation where {}. Turn it into a story.",
        "Write a story about what might happen if {}.",
        "Picture a world where {}. Write a story inspired by an r/AskReddit discussion.",
        "Imagine a person sharing an experience on r/AskReddit where {}. Tell their story.",

        # Creative and literary
        "Craft a fictional narrative centered on a situation where {}.",
        "Write an engaging tale inspired by an r/AskReddit thread in which {}.",
        "Develop a compelling story around the premise that {}.",
        "Use the scenario where {} as the foundation for a fictional narrative.",
        "Compose a short narrative depicting a situation in which {}.",

        # Informal and playful
        "Here's a wild r/AskReddit scenario: {}. Write a story about it.",
        "Imagine this happening: {}. Turn it into an entertaining story.",
        "Take this hypothetical situation—{}—and turn it into a story.",
        "What would a story look like if {}? Write one inspired by r/AskReddit.",
        "Suppose {}. Spin this scenario into a fictional tale.",
    ]

    def randomize_creative_en_prompt(prompt: str) -> str:
        """Reword recognized AskReddit story prefixes; leave other prompts unchanged."""
        about_match = re.match(
            r"^\s*Write a story from r/askreddit about\s+(.+?)\s*$",
            prompt,
            flags=re.IGNORECASE | re.DOTALL,
        )
        if about_match:
            return rng.choice(creative_en_story_about_templates).format(
                about_match.group(1).strip()
            )

        where_match = re.match(
            r"^\s*Write a story from r/askreddit where\s+(.+?)\s*$",
            prompt,
            flags=re.IGNORECASE | re.DOTALL,
        )
        if where_match:
            return rng.choice(creative_en_story_where_templates).format(
                where_match.group(1).strip()
            )

        # Avoid changing prompts from any other source or with another format.
        return prompt


    consumed_prefixes = set()

    ds_name = "agentlans/allenai-WildChat-4.8M"
    dataset = datasets.load_dataset(ds_name, "English", split="train", streaming=False).shuffle(seed=shuffle_seed)
    for sample in tqdm.tqdm(dataset, desc=ds_name+" English"):
        messages = sample["conversations"]
        if len(messages) == 2 and messages[0]["from"] == "human" and messages[1]["from"] == "gpt":
            q = messages[0]["value"].strip()
            a = messages[1]["value"].strip()
            if check_prompt_length(q) and 200 < len(a):
                # The WildChat "English" split contains some prompts in other languages.
                if is_english_text(q) and is_general_qa_sample(q, a):
                    prefix = q[:30]
                    if prefix not in consumed_prefixes:
                        consumed_prefixes.add(prefix)
                        prompts["qa"].append(q)
                        if len(prompts["qa"]) >= N_PREFETCH:
                            break

    ds_name = "smd20/social-engineering-qa-english"
    print(f"Loading {ds_name}...")
    ds = datasets.load_dataset(ds_name, split="train")
    for sample in ds.shuffle(seed=shuffle_seed).take(N_PREFETCH):
        q = sample["question"]
        a = sample["answer"]
        if is_english_text(q) and is_general_qa_sample(q, a):
            prefix = q[:30]
            if prefix not in consumed_prefixes:
                consumed_prefixes.add(prefix)
                prompts["qa"].append(q)

    # English creative writing
    ds_name = "IsDeeCee/StoryMaker"
    print(f"Loading {ds_name}...")
    ds = datasets.load_dataset(ds_name, split="train")
    for sample in ds.shuffle(seed=shuffle_seed):
        prompt = str(sample["Prompt"]).strip()
        prompts["creative-en"].append(randomize_creative_en_prompt(prompt))

    # Russian-English pairs
    print("Load pairs from Koziev/Translations...")
    URL = (
        "https://github.com/Koziev/Translations/raw/refs/heads/main/"
        "prose_translation_pairs.public.jsonl.gz"
    )
    for record in stream_jsonl_gz(URL):
        ru_text = record.get("right_text")
        if is_russian_text(ru_text) and len(ru_text) > 500:
            prompt = rng.choice(translation_templates).format(ru_text)
            if check_prompt_length(prompt):
                prompts["wmt ru-en"].append(prompt)

    ds_name = "/home/inkoziev/polygon/text_generator/tmp/sft_v053/sft23_dataset.unstressed.v053.eval.jsonl"
    print(f"Loading {ds_name}...")
    with open(ds_name) as f:
        for line in f:
            sample = json.loads(line)
            if sample["separator_token"] == "<poetry>":
                if len(sample["right_text"]) > 1000:
                    prompt = sample["left_text"]
                    prompts["poetry-ru"].append(prompt)
                    if len(prompts["poetry-ru"]) >= N_PREFETCH:
                        break

    # RUPOR dataset of GEC in different Russian text domains.
    print("Importing RUPOR...")
    src_fp = "/home/inkoziev/polygon/text_generator/py/generative_poetry/experiments/syntax_validation/Experiments4/datasets/RUPOR/rupor.json"
    gec_src_list = set()
    with open(src_fp) as f:
        data = json.load(f)
        for sample in data:
            if sample.get("corrupted_text") and sample.get("fixed_text") and sample["domain"] == "prose":
                gec_src_list.add(sample["corrupted_text"])

                if len(sample.get("fixed_text")) > 200:
                    prompts["wmt ru-en"].append(rng.choice(translation_templates).format(sample.get("fixed_text")))

    print("Importing LoRuGEC...")
    lorugec = build_lorugec()
    for sample in lorugec:
        gec_src_list.add(sample["text"])

    gec_src_list = sorted(list(gec_src_list), key=lambda text: -len(text))
    gec_src_list = gec_src_list[:N_PREFETCH]
    for src in gec_src_list:
        prompt = rng.choice(gec_ru_templates).format(src)
        prompts["gec-ru"].append(prompt)

    print("Importing JFLEG...")
    jfleg = build_jfleg()
    gec_src_list = set(sample["text"] for sample in jfleg)
    gec_src_list = sorted(list(gec_src_list), key=lambda text: -len(text))
    gec_src_list = gec_src_list[:N_PREFETCH]
    for src in gec_src_list:
        prompt = rng.choice(gec_en_templates).format(src)
        prompts["gec-en"].append(prompt)

    # Chinese poetry-generation instructions. Filter a broad Chinese instruction
    # corpus for prompts specifically about composing poetry or classical verse.
    # Randomly paraphrase the common prefix while preserving the imagery that follows.
    poetry_zh_original_prefix = "使用下列意象写一首现代诗："
    poetry_zh_prefix_templates = [
        "请根据以下意象创作一首现代诗：",
        "请运用下列意象写一首现代诗：",
        "以以下意象为灵感，写一首现代诗：",
        "请围绕下面这些意象创作一首现代诗：",
        "根据这些意象，创作一首现代诗：",
        "请用以下意象写一首现代诗：",
        "以这些意象为基础，写一首现代诗：",
        "请结合下列意象创作一首现代诗：",
        "请从以下意象出发，写一首现代诗：",
        "请将下列意象融入一首现代诗中：",
        "围绕以下意象，写一首现代诗：",
        "请选用下列意象来创作现代诗：",
        "请以这些意象为核心，创作一首现代诗：",
        "根据下面列出的意象写一首现代诗：",
        "请借助以下意象创作一首现代诗：",
        "把下列意象融入诗句，写成一首现代诗：",
        "请以以下意象为素材，写一首现代诗：",
        "请根据下面的意象写一首现代诗：",
        "请围绕这些意象写一首现代诗：",
        "运用下面给出的意象，创作一首现代诗：",
    ]

    def randomize_poetry_zh_prompt(prompt: str) -> str:
        """Replace the standard modern-poetry prefix, keeping its imagery unchanged."""
        if poetry_zh_original_prefix not in prompt:
            return prompt
        replacement_prefix = rng.choice(poetry_zh_prefix_templates)
        return prompt.replace(poetry_zh_original_prefix, replacement_prefix, 1)

    poetry_terms = ("写一首诗", "写首诗", "创作一首诗", "创作诗歌", "写一首古诗",
                    "作一首诗", "写诗", "诗歌创作", "填词", "写一首词", "对联",
                    "写一首歌", "写一首十四行诗", "诗")

    ds_name = "Iess/chinese_modern_poetry"
    print(f"Loading {ds_name} poetry candidates...")
    dataset = datasets.load_dataset(ds_name, split="train")
    for sample in tqdm.tqdm(dataset.shuffle(seed=shuffle_seed), desc=ds_name + " poetry"):
        q = str(sample.get("prompt", "")).strip()
        a = str(sample.get("response", "")).strip()
        if (
            q
            and a
            and is_chinese_text(q)
            and is_chinese_text(a)
            and any(term in q for term in poetry_terms)
            and check_chinese_prompt_length(q)
        ):
            prompts["poetry-zh"].append(randomize_poetry_zh_prompt(q))
            if len(prompts["poetry-zh"]) >= N_PREFETCH:
                break

    if len(prompts["poetry-zh"]) < N_PREFETCH:
        ds_name = "lcccher/Chinese-Instruct"
        print(f"Loading {ds_name} poetry candidates...")
        dataset = datasets.load_dataset(ds_name, "firefly", split="train")
        for sample in tqdm.tqdm(dataset.shuffle(seed=shuffle_seed), desc=ds_name + " poetry"):
            q = str(sample.get("prompt", "")).strip()
            a = str(sample.get("response", "")).strip()
            if (
                q
                and a
                and is_chinese_text(q)
                and is_chinese_text(a)
                and any(term in q for term in poetry_terms)
                and check_chinese_prompt_length(q)
            ):
                prompts["poetry-zh"].append(randomize_poetry_zh_prompt(q))
                if len(prompts["poetry-zh"]) >= N_PREFETCH:
                    break

    ds_name = "evilfreelancer/MATH-500-Russian"
    print(f"Loading {ds_name}...")
    ds = datasets.load_dataset(ds_name, split="test")
    for sample in ds.shuffle(seed=shuffle_seed):
        prompts["math-ru"].append(sample["problem"])

    ds_name = "agentlans/allenai-WildChat-4.8M"
    dataset = datasets.load_dataset(ds_name, "Russian", split="train", streaming=False).shuffle(seed=shuffle_seed)
    for sample in tqdm.tqdm(dataset, desc=ds_name+" Russian"):
        messages = sample["conversations"]
        if len(messages) == 2 and messages[0]["from"] == "human" and messages[1]["from"] == "gpt":
            q = messages[0]["value"].strip()
            prefix = q[:30]
            if prefix not in consumed_prefixes:
                consumed_prefixes.add(prefix)
                a = messages[1]["value"].strip()
                if check_prompt_length(q) and 200 < len(a) and is_russian_text(q) and is_russian_text(a):
                    prompts["qa-ru"].append(q)
                    if len(prompts["qa-ru"]) >= N_PREFETCH:
                        break

    # Chinese general question answering.
    ds_name = "agentlans/allenai-WildChat-4.8M"
    print(f"Loading {ds_name} Chinese...")
    dataset = datasets.load_dataset(ds_name, "Chinese", split="train", streaming=False).shuffle(seed=shuffle_seed)
    for sample in tqdm.tqdm(dataset, desc=ds_name + " Chinese"):
        messages = sample["conversations"]
        if len(messages) == 2 and messages[0]["from"] == "human" and messages[1]["from"] == "gpt":
            q = messages[0]["value"].strip()
            a = messages[1]["value"].strip()
            if (
                200 < len(q)
                and 200 < len(a)
                and is_chinese_text(q)
                and is_chinese_text(a)
                and check_chinese_prompt_length(q)
            ):
                prompts["qa-zh"].append(q)
                if len(prompts["qa-zh"]) >= N_PREFETCH:
                    break

    # Chinese mathematics prompt templates. Randomly vary the instruction prefix
    # while preserving the problem statement verbatim.
    math_zh_original_prefix = "请详细解答下面的数学问题，给出清晰的推导步骤和最终答案："
    math_zh_prefix_templates = [
        "请详细解答以下数学问题，并写出清晰的推导过程和最终答案：",
        "请解答下面这道数学题，列出明确的推理步骤并给出最终答案：",
        "请逐步分析并解答以下数学问题，说明推导过程和答案：",
        "请给出下列数学问题的详细解法、推导步骤及最终结果：",
        "请认真解答以下数学题，展示清楚的计算过程并写出答案：",
        "请对下面的数学问题进行详细推导，并给出最后的答案：",
        "请一步一步解答下列数学问题，说明思路、推导过程和最终答案：",
        "请分析下面这道数学题，给出完整的解题步骤和最终结果：",
        "请详细说明以下数学问题的求解过程，并写出最终答案：",
        "请展示解答下列数学问题的清晰步骤，并给出答案：",
        "请对以下数学题进行逐步求解，列明推导过程和最终答案：",
        "请写出下面数学问题的完整解答，包括推理步骤和最后结果：",
        "请清楚地推导并解答以下问题，最后给出数学答案：",
        "请解出下面的数学问题，并提供详细的计算或证明步骤：",
        "请按步骤求解以下数学问题，解释推导过程并给出最终结果：",
        "请为下面这道数学题提供详细解答，包含清晰的步骤和答案：",
        "请说明如何求解下列数学问题，逐步推导并写出最终答案：",
        "请完整解答以下数学问题，确保推导过程清楚、答案明确：",
        "请先展示下面数学问题的解题思路和推导步骤，再给出答案：",
        "请针对以下数学问题给出分步解答，并明确写出最终答案：",
    ]

    def randomize_math_zh_prompt(prompt):
        """Randomize the standard Chinese math instruction prefix only."""
        if math_zh_original_prefix not in prompt:
            return prompt
        return prompt.replace(math_zh_original_prefix, rng.choice(math_zh_prefix_templates), 1)

    def canonicalize_math_zh_prompt(prompt):
        """Normalize any known math prefix for duplicate detection."""
        for prefix in math_zh_prefix_templates:
            if prefix in prompt:
                return prompt.replace(prefix, math_zh_original_prefix, 1)
        return prompt

    # Chinese mathematics. LiveK12Bench provides Chinese and English splits;
    # use only Chinese splits and retain math questions.
    ds_name = "Shawn-wxh/livek12bench"
    print(f"Loading {ds_name} Chinese mathematics...")
    livek12 = datasets.load_dataset(ds_name).shuffle(seed=shuffle_seed)
    for split_name in ("zh_2603", "zh_2605"):
        if split_name not in livek12:
            continue
        for sample in tqdm.tqdm(livek12[split_name], desc=f"{ds_name}:{split_name}"):
            if sample.get("subject") != "math":
                continue
            question = str(sample.get("question", "")).strip()
            if question and is_chinese_text(question):
                prompt = math_zh_original_prefix + "\n\n" + question
                if check_prompt_length(prompt):
                    prompts["math-zh"].append(randomize_math_zh_prompt(prompt))
                    if len(prompts["math-zh"]) >= N_PREFETCH:
                        break
        if len(prompts["math-zh"]) >= N_PREFETCH:
            break


    # Supplement Chinese mathematics with native Chinese word problems
    # from CMATH, then translated GSM8K questions if more candidates are needed.
    math_zh_seen = {canonicalize_math_zh_prompt(prompt) for prompt in prompts["math-zh"]}

    def add_math_zh_question(question):
        question = str(question or "").strip()
        if not question or not is_chinese_text(question):
            return

        prompt = math_zh_original_prefix + "\n\n" + question

        if prompt in math_zh_seen or not check_prompt_length(prompt):
            return

        prompts["math-zh"].append(randomize_math_zh_prompt(prompt))
        math_zh_seen.add(prompt)

    # 1. Native Chinese math word problems.
    if len(prompts["math-zh"]) < N_PREFETCH:
        ds_name = "weitianwen/cmath"
        print(f"Loading {ds_name}...")
        cmath = datasets.load_dataset(ds_name, split="validation")

        for sample in cmath.shuffle(seed=shuffle_seed):
            add_math_zh_question(sample.get("question", ""))

            if len(prompts["math-zh"]) >= N_PREFETCH:
                break

    # 2. Additional Chinese-translated grade-school math problems.
    if len(prompts["math-zh"]) < N_PREFETCH:
        ds_name = "meta-math/GSM8K_zh"
        print(f"Loading {ds_name}...")
        gsm8k_zh = datasets.load_dataset(ds_name, split="train")

        for sample in gsm8k_zh.shuffle(seed=shuffle_seed):
            add_math_zh_question(sample.get("question_zh", ""))

            if len(prompts["math-zh"]) >= N_PREFETCH:
                break

    print(f"math-zh candidate prompts after supplementation: "
          f"{len(prompts['math-zh'])}")


    # Chinese coding instructions. The "code" configuration contains prompt/response pairs.
    ds_name = "Mxode/Chinese-Instruct-Lite"
    print(f"Loading {ds_name} code...")
    dataset = datasets.load_dataset(ds_name, "code", split="train")
    for sample in tqdm.tqdm(dataset.shuffle(seed=shuffle_seed), desc=ds_name + " code"):
        q = str(sample.get("prompt", "")).strip()
        a = str(sample.get("response", "")).strip()
        if (
            q
            and a
            and is_chinese_text(q)
            and check_prompt_length(q)
        ):
            prompts["code-zh"].append(q)
            if len(prompts["code-zh"]) >= N_PREFETCH:
                break

    # MERA-evaluation/MERA
    ds_name = "MERA-evaluation/MERA"
    print(f"Loading {ds_name}...")
    configs = datasets.get_dataset_config_names(ds_name)
    n = 0
    for config in configs:
        for split_name in ["test", "validation"]:
            try:
                splits = datasets.get_dataset_split_names(ds_name, config)
                if split_name not in splits:
                    continue

                dataset = datasets.load_dataset(ds_name, config, split=split_name)
                for sample in tqdm.tqdm(dataset, desc=ds_name + ":" + config + ":" + split_name):
                    instruct = sample['instruction']
                    if isinstance(sample['inputs'], str):
                        instruct = re.sub(r'\{.+?\}', sample['inputs'], instruct)
                    else:
                        for key, value in sample['inputs'].items():
                            instruct = instruct.replace('{' + key + '}', str(value))

                    if config == 'multiq':
                        output = sample['outputs']
                    elif config == "ruhumaneval":
                        if 'function' in sample['inputs']:
                            output = sample['inputs']['function'] + sample['meta']["canonical_solution"]
                        else:
                            output = sample['meta']["canonical_solution"]
                    elif isinstance(sample['outputs'], str):
                        output = sample['outputs'].strip()
                    else:
                        raise NotImplementedError()

                    q = instruct
                    a = output
                    if check_prompt_length(q) and a:
                        if config in ["ruhumaneval", "rucodeeval"]:
                            prompts["code-ru"].append(q)
                        else:
                            prompts["qa-ru"].append(q)

            except Exception as ex:
                print(traceback.format_exc())

    if len(prompts["code-ru"]) < N_PREFETCH:
        with open("coding_data/ru_coding_instruct_dataset.jsonl") as f:
            for line in f:
                sample = json.loads(line)
                prompt = sample["query"]
                if check_prompt_length(prompt) and is_russian_text(prompt[:100]):
                    prompts["code-ru"].append(prompt)

    ds_name = "wmt/wmt19"
    print(f"Loading {ds_name}...")
    dataset = datasets.load_dataset(ds_name, "ru-en", split="validation")

    ru_en_pairs = []
    for example in dataset:
        src = example["translation"]["ru"]
        ref = example["translation"]["en"]
        # отбираем тексты не слишком короткие (в wmt полно однословных сэмплов).
        if check_prompt_length(src) and 200 < len(ref) < 2000 and is_russian_text(src):
            prompts["wmt ru-en"].append(f"Translate the text into English: {src}")

    ds_name = "me-aas/python-code-dataset-500k"
    print(f"Loading {ds_name}...")
    ds = datasets.load_dataset(ds_name, split="train")
    for sample in ds.shuffle(seed=shuffle_seed).take(N_PREFETCH):
        prompts["code"].append(sample["instruction"])

    ds_name = "HaimingW/math_train_decontaminated"
    print(f"Loading {ds_name}...")
    ds = datasets.load_dataset(ds_name, split="train")
    for sample in ds.shuffle(seed=shuffle_seed).take(N_PREFETCH):
        prompts["math"].append(sample["problem"])

    ds_name = "checkai/instruction-poems"
    print(f"Loading {ds_name}...")
    ds = datasets.load_dataset(ds_name, split="train")
    for sample in ds.shuffle(seed=shuffle_seed).take(N_PREFETCH):
        prompts["poetry-en"].append(sample["INSTRUCTION"])

    # ---------------------------------------------------------------------
    # Let’s make a final selection of prompts that will elicit the longest
    # possible responses from the model (batched greedy generation).
    final_prompts = collections.defaultdict(list)

    # Per-domain properties of prompts and corresponding responses
    qa_prompts = collections.defaultdict(list)
    math_prompts = collections.defaultdict(list)
    code_prompts = collections.defaultdict(list)
    poetry_prompts = collections.defaultdict(list)

    print(f"Loading {model_id}...")
    model = transformers.AutoModelForCausalLM.from_pretrained(
        model_id,
        dtype=torch.bfloat16, device_map="cuda",
        attn_implementation="flash_attention_2",  # options: sdpa | eager | flash_attention_2 | flash_attention_4
        trust_remote_code=True,
    ).eval()

    info_table = [["Domain", "No. prompts", "Prompt length, toks<br>mean<sub>min, max</sub>", "Response length, toks<br>mean<sub>min, max</sub>"]]
    for idomain, domain in enumerate(sorted(prompts.keys()), start=1):
        print("\n{}/{} Start processing domain \"{}\"".format(idomain, len(prompts), domain))

        raw_prompts = prompts[domain]
        raw_prompts = list(dict.fromkeys(raw_prompts))  # avoid duplicates; keeps the order, so the selection is reproducible
        if domain.endswith("-zh"):
            raw_prompts = [prompt for prompt in raw_prompts if check_chinese_prompt_length(prompt)]
        else:
            raw_prompts = [prompt for prompt in raw_prompts if check_prompt_length(prompt)]

        # The WildChat "English" split contains some prompts in other languages.
        # Filter the complete qa candidate pool here, after merging all sources
        # and removing duplicates, so every prompt entering the English qa domain
        # is checked (including prompts added from the supplementary QA dataset).
        if domain == "qa":
            before_language_filter = len(raw_prompts)
            raw_prompts = [prompt for prompt in raw_prompts if is_english_text(prompt)]
            rejected = before_language_filter - len(raw_prompts)
            print(
                f"English language filter for qa: kept {len(raw_prompts)}, "
                f"rejected {rejected} non-English/undetermined prompts."
            )

        raw_prompts = sorted(raw_prompts, key=lambda z: -len(z))
        print("Raw prompts: {}".format(len(raw_prompts)))

        if not raw_prompts:
            print(f"WARNING: no candidate prompts collected for domain {domain}; skipping.")
            continue

        if len(raw_prompts) > N_REPEATS*4:
            raw_prompts = raw_prompts[:N_REPEATS*4]
            print("No. raw prompts after truncation: {}".format(len(raw_prompts)))

        prompt_response_list = measure_response_lengths(model, tokenizer, raw_prompts, GEN_BATCH_SIZE, MAX_NEW_TOKENS,
                                                        desc=domain)

        prompt_response_list = sorted(prompt_response_list, key=lambda z: -z[2])  # stable: ties keep the original order
        prompt_response_list = prompt_response_list[:N_REPEATS]

        if domain.endswith("-ru"):
            language = "ru"
        elif domain.endswith("-zh"):
            language = "zh"
        else:
            language = "en"

        for prompt, prompt_num_toks, response_num_toks in prompt_response_list:
            final_prompts[domain].append(prompt)
            prompt_info = {"prompt": prompt,
                           "prompt_num_toks": prompt_num_toks,
                           "response_num_toks": response_num_toks}
            if domain.startswith("math"):
                math_prompts[language].append(prompt_info)
            elif domain.startswith("qa"):
                qa_prompts[language].append(prompt_info)
            elif domain.startswith("code"):
                code_prompts[language].append(prompt_info)
            elif domain.startswith("poetry"):
                poetry_prompts[language].append(prompt_info)

        prompt_lengths = [prompt_len for _, prompt_len, _ in prompt_response_list]
        response_lengths = [response_len for _, _, response_len in prompt_response_list]
        if not response_lengths:
            print(f"WARNING: no measured prompts for domain {domain}; skipping statistics.")
            final_prompts.pop(domain, None)
            continue

        print("Domain: {}  population: {}  response lengths: min={} avg={} max={}".format(domain, len(final_prompts[domain]), min(response_lengths), np.mean(response_lengths), max(response_lengths)))

        info_table.append((domain,
                           len(final_prompts[domain]),
                           "{:.1f}<sub>{}, {}</sub>".format(np.mean(prompt_lengths), min(prompt_lengths), max(prompt_lengths)),
                           "{:.1f}<sub>{}, {}</sub>".format(np.mean(response_lengths), min(response_lengths), max(response_lengths))
                           ))

        # After processing each domain, we output the selected prompts to a file for visual review.
        with open("speed_evaluation_dataset.txt", "w") as f:
            for domain, domain_prompts in final_prompts.items():
                f.write(f"\n=== START OF DOMAIN: {domain} ===\n\n")
                for i, prompt in enumerate(domain_prompts, start=1):
                    f.write(f"--- domain: {domain}  prompt #{i} ---\n\n{prompt}\n\n")
                f.write(f"\n=== END OF DOMAIN {domain} ===\n\n")

        # Visualize prompt and response lengths for several domains
        for domain, stat in [("code", code_prompts), ("qa", qa_prompts), ("math", math_prompts), ("poetry", poetry_prompts)]:
            if stat:
                language_to_prompt_lengths = dict()
                for language, stat in stat.items():
                    language_to_prompt_lengths[language] = [x["response_num_toks"] for x in stat]

                fig, zx = plot_grouped_length_histograms(domain, language_to_prompt_lengths, bins=5, density=True)
                try:
                    fig.savefig(f"response_length_in_domain_{domain}.png", dpi=150, bbox_inches="tight")
                finally:
                    plt.close(fig)

        # Generation speed evaluation dataset in JSON file.
        with open("speed_evaluation_dataset.json", "w") as f:
            json.dump(final_prompts, f, indent=4, ensure_ascii=False)

        # Dataset statistics in Markdown format.
        with open("speed_evaluation_dataset.md", "w") as f:
            f.write("Tokenizer: `{}`\n\n".format(model_id))
            f.write(terminaltables.GithubFlavoredMarkdownTable(info_table).table)
